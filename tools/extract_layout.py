"""Extract a compact print-layout spec (widths, heights, styles, static text,
merges) from the two Evertz templates, so the browser can draw a faithful PDF
version of each weekly sheet with jsPDF."""
import json, os, re, colorsys, sys, zipfile
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string, range_boundaries

THEMES = {}

def theme_palette(path):
    z = zipfile.ZipFile(path)
    xml = z.read('xl/theme/theme1.xml').decode('utf8')
    scheme = re.search(r'<a:clrScheme.*?</a:clrScheme>', xml, re.S).group(0)
    def grab(tag):
        m = re.search(r'<a:%s>(.*?)</a:%s>' % (tag, tag), scheme, re.S).group(1)
        v = re.search(r'(?:val|lastClr)="([0-9A-Fa-f]{6})"', m.replace('val="window"', '').replace('val="windowText"', ''))
        return v.group(1).upper()
    dk1, lt1, dk2, lt2 = grab('dk1'), grab('lt1'), grab('dk2'), grab('lt2')
    acc = [grab('accent%d' % i) for i in range(1, 7)]
    # Excel's theme index swaps the light/dark pairs
    return [lt1, dk1, lt2, dk2] + acc + [grab('hlink'), grab('folHlink')]

def apply_tint(hexrgb, tint):
    if not tint:
        return hexrgb
    r, g, b = [int(hexrgb[i:i+2], 16) / 255 for i in (0, 2, 4)]
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, l, s)
    return '%02X%02X%02X' % (round(r * 255), round(g * 255), round(b * 255))

def color(c, pal, default=None):
    if c is None:
        return default
    try:
        if c.type == 'rgb' and isinstance(c.rgb, str):
            v = c.rgb[-6:].upper()
            return v
        if c.type == 'theme':
            return apply_tint(pal[c.theme], c.tint)
        if c.type == 'indexed':
            idx = c.indexed
            if idx in (64, 65):
                return default
            from openpyxl.styles.colors import COLOR_INDEX
            return COLOR_INDEX[idx][-6:].upper()
    except Exception:
        return default
    return default

def col_widths_from_xml(path, sheet_file, ncols):
    z = zipfile.ZipFile(path)
    xml = z.read('xl/worksheets/' + sheet_file).decode('utf8')
    dflt = float(re.search(r'defaultColWidth="([\d.]+)"', xml).group(1))
    widths = [dflt] * (ncols + 1)
    for m in re.finditer(r'<col ([^>]*)/>', xml):
        a = dict(re.findall(r'(\w+)="([^"]*)"', m.group(1)))
        lo, hi, w = int(a['min']), int(a['max']), float(a['width'])
        hidden = a.get('hidden') == '1'
        for i in range(lo, min(hi, ncols) + 1):
            widths[i] = 0 if hidden else w
    dflt_h = float(re.search(r'defaultRowHeight="([\d.]+)"', xml).group(1))
    margins = dict(re.findall(r'(left|right|top|bottom)="([\d.]+)"', re.search(r'<pageMargins[^>]*>', xml).group(0)))
    return widths[1:], dflt_h, {k: float(v) * 72 for k, v in margins.items()}

def char_to_pt(w, mdw):
    # Excel column width is in units of the Normal font's max digit width (px @96dpi)
    if w <= 0:
        return 0
    px = int(w * mdw + 5) if w >= 1 else int(w * (mdw + 5) + 0.5)
    return round(px * 0.75, 2)

def extract(path, sheet_name, sheet_file, rng, dynamic, mdw):
    pal = theme_palette(path)
    wb = openpyxl.load_workbook(path)
    wbv = openpyxl.load_workbook(path, data_only=True)
    ws, wsv = wb[sheet_name], wbv[sheet_name]
    c1, r1, c2, r2 = range_boundaries(rng)
    widths, dflt_h, margins = col_widths_from_xml(path, sheet_file, c2)
    cols = [char_to_pt(w, mdw) for w in widths[c1 - 1:c2]]
    rows = []
    for r in range(r1, r2 + 1):
        d = ws.row_dimensions.get(r)
        rows.append(d.height if d is not None and d.height is not None else dflt_h)
    styles, style_ix, cells = [], {}, []
    dyn = set()
    for d in dynamic:
        if ':' in d:
            a, b, cc, dd = range_boundaries(d)
            for rr in range(b, dd + 1):
                for c in range(a, cc + 1):
                    dyn.add(get_column_letter(c) + str(rr))
        else:
            dyn.add(d)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ref = get_column_letter(c) + str(r)
            cell = ws[ref]
            f = cell.font
            fill = None
            if cell.fill is not None and cell.fill.fill_type == 'solid':
                fill = color(cell.fill.fgColor, pal)
            bd = cell.border
            def side(s):
                if s is None or s.style is None:
                    return None
                return [s.style, color(s.color, pal, '000000')]
            al = cell.alignment
            st = {
                'fn': (f.name or 'Arial'), 'sz': float(f.sz or 10), 'b': bool(f.b), 'i': bool(f.i),
                'u': bool(f.u),
                'fc': color(f.color, pal, '000000'), 'fill': fill,
                'bl': side(bd.left), 'br': side(bd.right), 'bt': side(bd.top), 'bb': side(bd.bottom),
                'ha': al.horizontal, 'va': al.vertical, 'wrap': bool(al.wrap_text),
                'nf': cell.number_format,
            }
            v = cell.value
            out_v = None
            if ref not in dyn and v is not None:
                if isinstance(v, str) and v.startswith('='):
                    cv = wsv[ref].value
                    out_v = cv
                else:
                    out_v = v
            if hasattr(out_v, 'isoformat'):
                # Excel serial date
                import datetime
                base = datetime.datetime(1899, 12, 30)
                out_v = {'d': (out_v - base).days}
            has_vis = fill or any(st[k] for k in ('bl', 'br', 'bt', 'bb'))
            if not has_vis and out_v is None and ref not in dyn:
                continue
            key = json.dumps(st, sort_keys=True)
            if key not in style_ix:
                style_ix[key] = len(styles)
                styles.append(st)
            entry = [ref, style_ix[key]]
            if out_v is not None:
                entry.append(out_v)
            cells.append(entry)
    merges = [str(m) for m in ws.merged_cells.ranges]
    # keep merges inside range
    def inside(m):
        a, b, cc, dd = range_boundaries(m)
        return a >= c1 and b >= r1 and cc <= c2 and dd <= r2
    merges = [m for m in merges if inside(m)]
    return {
        'range': rng, 'cols': cols, 'rows': rows, 'margins': margins,
        'styles': styles, 'cells': cells, 'merges': merges,
    }

itemized = extract(
    os.path.join(ROOT, 'templates', 'itemized-expense-claim.xlsx'), 'Week Ending 02-May-26', 'sheet1.xml', 'A1:Q40',
    ['N2', 'N3', 'N4', 'N5', 'B10:O10', 'B11:O12', 'B14:Q31', 'B32:Q32', 'F34'], 9)
perdiem = extract(
    os.path.join(ROOT, 'templates', 'perdiem-travel-expense.xlsx'), 'Expense report', 'sheet1.xml', 'A1:W41',
    ['N4', 'N5', 'N6', 'N7', 'N8', 'B15:W29', 'G30:T30', 'U31:W31', 'U33'], 7)
out = {'itemized': itemized, 'perdiem': perdiem}
s = json.dumps(out, separators=(',', ':'))
open(os.path.join(ROOT, 'src', 'layouts.json'), 'w').write(s)
print(len(s), 'bytes;', len(itemized['cells']), 'itemized cells,', len(perdiem['cells']), 'perdiem cells;',
      len(itemized['styles']), '/', len(perdiem['styles']), 'styles')
print('itemized cols', itemized['cols'], 'sum', sum(itemized['cols']), 'rows sum', sum(itemized['rows']))
print('perdiem cols', perdiem['cols'], 'sum', sum(perdiem['cols']), 'rows sum', sum(perdiem['rows']))
print('margins', itemized['margins'], perdiem['margins'])
print(sorted(set(s['nf'] for s in itemized['styles'] + perdiem['styles'])))
print(sorted(set(s['fn'] for s in itemized['styles'] + perdiem['styles'])))
