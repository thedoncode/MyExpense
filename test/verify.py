"""Checks the workbooks in test/out (run `node test/generate.js` first).

1. Template fidelity: on every weekly sheet, everything outside the data cells (labels, formulas,
   totals, the per diem calculator) is exactly the template's, and the per diem reference tabs are
   unchanged.
2. Snapshots: every cell of every weekly sheet (formula and value) matches test/golden/. The snapshots
   were recorded from the generator after it had reproduced the original filled-in Evertz workbooks
   cell for cell. After an intended change, re-record them with `python3 test/verify.py --update`.
3. LibreOffice recalculates every workbook with zero formula errors, and its results match the cached
   values MyExpense wrote.

Needs: python3 with openpyxl, and LibreOffice (`soffice`) for step 3.
"""
import datetime, glob, json, os, re, shutil, subprocess, sys, tempfile, warnings
from pathlib import Path
import openpyxl
from openpyxl.utils import range_boundaries, get_column_letter

warnings.filterwarnings('ignore')
ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'test' / 'out'
GOLDEN = ROOT / 'test' / 'golden'
UPDATE = '--update' in sys.argv
failures = 0

def refs(*ranges):
    out = set()
    for rg in ranges:
        c1, r1, c2, r2 = range_boundaries(rg if ':' in rg else rg + ':' + rg)
        out |= {get_column_letter(c) + str(r) for c in range(c1, c2 + 1) for r in range(r1, r2 + 1)}
    return out

KINDS = {
    'itemized': {'template': 'itemized-expense-claim.xlsx', 'base': 0,
                 'data': refs('N2:N5', 'B10:O10', 'B12:O12', 'B14:O31')},
    'perdiem': {'template': 'perdiem-travel-expense.xlsx', 'base': 0,
                'data': refs('N4:N8', 'B15:W29')},
}

def plain(v):
    if isinstance(v, datetime.datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, float):
        return round(v, 6)
    return v

def snapshot(path):
    f, v = openpyxl.load_workbook(path), openpyxl.load_workbook(path, data_only=True)
    snap = {}
    for ws in f.worksheets:
        if not ws.title.startswith('Week Ending'):
            continue
        snap[ws.title] = {c.coordinate: [plain(c.value), plain(v[ws.title][c.coordinate].value)]
                          for row in ws.iter_rows() for c in row if c.value is not None}
    return snap

for path in sorted(glob.glob(str(OUT / '*.xlsx'))):
    name = os.path.basename(path)
    tag = name.split(' - ')[0]
    kind = KINDS['perdiem' if tag.startswith('pd') else 'itemized']
    tpl = openpyxl.load_workbook(ROOT / 'templates' / kind['template'])
    base = tpl.worksheets[kind['base']]
    out = openpyxl.load_workbook(path)

    # 1. template fidelity
    bad = []
    for ws in out.worksheets:
        if ws.title.startswith('Week Ending'):
            cells = {c.coordinate for row in base.iter_rows() for c in row} | {c.coordinate for row in ws.iter_rows() for c in row}
            bad += [(ws.title, r, base[r].value, ws[r].value) for r in cells - kind['data'] if base[r].value != ws[r].value]
        else:
            ref = tpl[ws.title]
            cells = {c.coordinate for row in ref.iter_rows() for c in row} | {c.coordinate for row in ws.iter_rows() for c in row}
            bad += [(ws.title, r, ref[r].value, ws[r].value) for r in cells if ref[r].value != ws[r].value]
    print(f'{name}: template differences outside data cells: {len(bad)}')
    for b in bad[:10]:
        print('   ', b)
    failures += len(bad)

    # 2. snapshots
    snap = snapshot(path)
    gpath = GOLDEN / (tag + '.json')
    if UPDATE:
        GOLDEN.mkdir(exist_ok=True)
        gpath.write_text(json.dumps(snap, indent=1, sort_keys=True))
        print(f'    snapshot recorded: {gpath.name}')
    else:
        want = json.loads(gpath.read_text())
        diffs = [(s, r, want.get(s, {}).get(r), snap.get(s, {}).get(r))
                 for s in set(want) | set(snap) for r in set(want.get(s, {})) | set(snap.get(s, {}))
                 if want.get(s, {}).get(r) != snap.get(s, {}).get(r)]
        print(f'    snapshot differences: {len(diffs)}')
        for d in diffs[:10]:
            print('   ', d)
        failures += len(diffs)

MACRO = '''<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
Sub Recalc()
  ThisComponent.calculateAll()
  ThisComponent.store()
  ThisComponent.close(True)
End Sub
</script:module>'''

def recalc(path):
    with tempfile.TemporaryDirectory() as prof:
        url = Path(prof).as_uri()
        subprocess.run(['soffice', '--headless', '--terminate_after_init', f'-env:UserInstallation={url}'], capture_output=True, timeout=120)
        (Path(prof) / 'user' / 'basic' / 'Standard' / 'Module1.xba').write_text(MACRO)
        subprocess.run(['soffice', '--headless', f'-env:UserInstallation={url}', 'macro:///Standard.Module1.Recalc', str(path)], capture_output=True, timeout=180)

# 3. LibreOffice recalculation
if shutil.which('soffice'):
    for f in sorted(glob.glob(str(OUT / '*.xlsx'))):
        with tempfile.TemporaryDirectory() as tmp:
            copy = Path(tmp) / 'copy.xlsx'
            shutil.copy(f, copy)
            recalc(copy)
            mine, lo = openpyxl.load_workbook(f, data_only=True), openpyxl.load_workbook(copy, data_only=True)
            errs = mism = 0
            for ws in mine.worksheets:
                for row in lo[ws.title].iter_rows():
                    for c in row:
                        v1, v2 = ws[c.coordinate].value, c.value
                        if isinstance(v2, str) and v2.startswith('#'):
                            errs += 1
                        norm = lambda v: None if v in ('', None) else (round(v, 6) if isinstance(v, (int, float)) else v)
                        if norm(v1) != norm(v2):
                            mism += 1
            print(f'{os.path.basename(f)}: formula errors {errs}, cached-value mismatches {mism}')
            failures += errs + mism
else:
    print('soffice not found: skipped the LibreOffice recalculation check')

print('PASS' if failures == 0 else f'FAIL ({failures})')
sys.exit(1 if failures else 0)
